"""
Generate ZavaTax architecture diagrams in SVG and PPTX formats.
All diagrams use 16:9 aspect ratio (1920x1080 for SVG, standard slide for PPTX).
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.dml.color import RGBColor
import os

# =============================================================================
# Color Palette (Accessible, high contrast)
# =============================================================================
COLORS = {
    'web_app': '2563EB',       # Blue - Web Application
    'web_app_dark': '1D4ED8',
    'dab': '059669',           # Emerald - Data API Builder
    'dab_dark': '047857',
    'primary': '7C3AED',       # Purple - Primary Replica
    'primary_dark': '6D28D9',
    'ha_replica': '8B5CF6',    # Lighter purple - HA Replica
    'named_replica': '0891B2', # Cyan - Named Replica
    'azure_ai': 'EA580C',      # Orange - Azure OpenAI
    'azure_ai_dark': 'C2410C',
    'arrow_rw': '16A34A',      # Green - Read/Write
    'arrow_read': 'F59E0B',    # Amber - Read-only
    'arrow_analytics': '0891B2', # Cyan - Analytics
    'arrow_ai': 'DC2626',      # Red - AI calls
    'text_dark': '1E293B',     # Dark text
    'text_light': 'FFFFFF',    # White text
    'border': '94A3B8',        # Gray border
    'bg': 'F8FAFC',            # Light background
    'hyperscale_border': '7C3AED',
    'ai_border': 'EA580C',
}

def hex_to_rgb(hex_color):
    """Convert hex color to RGBColor."""
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))

# =============================================================================
# SVG Generation
# =============================================================================

def generate_svg_diagram1():
    """Diagram 1: Base Architecture - Web App + DAB + Hyperscale (Primary + HA)"""
    return '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1920 1080" font-family="Segoe UI, Arial, sans-serif">
  <defs>
    <filter id="shadow" x="-4%" y="-4%" width="108%" height="112%">
      <feDropShadow dx="3" dy="4" stdDeviation="6" flood-opacity="0.12"/>
    </filter>
    <marker id="arrowGreen" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#16A34A"/>
    </marker>
    <marker id="arrowAmber" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#F59E0B"/>
    </marker>
    <marker id="arrowGray" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#64748B"/>
    </marker>
  </defs>

  <!-- Background -->
  <rect width="1920" height="1080" fill="#F8FAFC"/>

  <!-- Title -->
  <text x="960" y="70" text-anchor="middle" font-size="36" font-weight="700" fill="#1E293B">ZavaTax Architecture — Base</text>
  <text x="960" y="110" text-anchor="middle" font-size="20" fill="#64748B">Web Application + Data API Builder + Azure SQL DB Hyperscale</text>

  <!-- ═══════════════════ WEB APP ═══════════════════ -->
  <rect x="80" y="280" width="340" height="400" rx="20" fill="#2563EB" filter="url(#shadow)"/>
  <text x="250" y="340" text-anchor="middle" font-size="26" font-weight="700" fill="#FFF">Web Application</text>
  <line x1="120" y1="360" x2="380" y2="360" stroke="#93C5FD" stroke-width="2"/>
  <text x="250" y="400" text-anchor="middle" font-size="18" fill="#DBEAFE">React + TypeScript</text>
  <text x="250" y="430" text-anchor="middle" font-size="18" fill="#DBEAFE">TailwindCSS</text>
  
  <rect x="110" y="470" width="280" height="50" rx="10" fill="rgba(255,255,255,0.15)"/>
  <text x="250" y="502" text-anchor="middle" font-size="18" font-weight="600" fill="#FFF">CQRS Workload Router</text>
  
  <rect x="110" y="540" width="280" height="110" rx="10" fill="rgba(255,255,255,0.1)"/>
  <text x="250" y="575" text-anchor="middle" font-size="16" fill="#DBEAFE">Routes requests by workload type</text>
  <text x="250" y="600" text-anchor="middle" font-size="16" fill="#DBEAFE">• Read/Write → Primary</text>
  <text x="250" y="625" text-anchor="middle" font-size="16" fill="#DBEAFE">• Read-Only → HA Replica</text>

  <!-- ═══════════════════ DAB ═══════════════════ -->
  <rect x="580" y="280" width="340" height="400" rx="20" fill="#059669" filter="url(#shadow)"/>
  <text x="750" y="340" text-anchor="middle" font-size="26" font-weight="700" fill="#FFF">Data API Builder</text>
  <line x1="620" y1="360" x2="880" y2="360" stroke="#6EE7B7" stroke-width="2"/>
  <text x="750" y="400" text-anchor="middle" font-size="18" fill="#D1FAE5">REST + GraphQL API</text>
  
  <rect x="610" y="440" width="280" height="70" rx="10" fill="rgba(255,255,255,0.2)"/>
  <circle cx="640" cy="475" r="12" fill="#FDE047"/>
  <text x="665" y="465" font-size="16" font-weight="600" fill="#FFF">Primary Data Source</text>
  <text x="665" y="490" font-size="14" fill="#D1FAE5">Read/Write operations</text>

  <rect x="610" y="530" width="280" height="70" rx="10" fill="rgba(255,255,255,0.2)"/>
  <circle cx="640" cy="565" r="12" fill="#F59E0B"/>
  <text x="665" y="555" font-size="16" font-weight="600" fill="#FFF">Read-Only Data Source</text>
  <text x="665" y="580" font-size="14" fill="#D1FAE5">ApplicationIntent=ReadOnly</text>

  <rect x="610" y="620" width="280" height="40" rx="8" fill="rgba(0,0,0,0.15)"/>
  <text x="750" y="646" text-anchor="middle" font-size="14" fill="#FFF">Auto-generated CRUD + Stored Procs</text>

  <!-- ═══════════════════ HYPERSCALE ═══════════════════ -->
  <rect x="1080" y="200" width="760" height="580" rx="24" fill="none" stroke="#7C3AED" stroke-width="4" stroke-dasharray="12 6"/>
  <text x="1460" y="250" text-anchor="middle" font-size="24" font-weight="700" fill="#7C3AED">Azure SQL DB Hyperscale</text>
  <text x="1460" y="280" text-anchor="middle" font-size="16" fill="#A78BFA">Zone-Redundant High Availability</text>

  <!-- Primary Replica -->
  <rect x="1130" y="320" width="320" height="200" rx="16" fill="#7C3AED" filter="url(#shadow)"/>
  <text x="1290" y="375" text-anchor="middle" font-size="24" font-weight="700" fill="#FFF">Primary Replica</text>
  <text x="1290" y="410" text-anchor="middle" font-size="18" fill="#EDE9FE">Read / Write</text>
  <rect x="1160" y="435" width="260" height="60" rx="8" fill="rgba(255,255,255,0.15)"/>
  <text x="1290" y="462" text-anchor="middle" font-size="14" fill="#EDE9FE">Tables · Stored Procedures</text>
  <text x="1290" y="482" text-anchor="middle" font-size="14" fill="#EDE9FE">VECTOR · DiskANN · Columnstore</text>

  <!-- HA Replica -->
  <rect x="1500" y="320" width="320" height="200" rx="16" fill="#8B5CF6" filter="url(#shadow)"/>
  <text x="1660" y="375" text-anchor="middle" font-size="24" font-weight="700" fill="#FFF">HA Replica</text>
  <text x="1660" y="410" text-anchor="middle" font-size="18" fill="#EDE9FE">Read-Only</text>
  <rect x="1530" y="435" width="260" height="60" rx="8" fill="rgba(255,255,255,0.15)"/>
  <text x="1660" y="462" text-anchor="middle" font-size="14" fill="#EDE9FE">Zone-Redundant</text>
  <text x="1660" y="482" text-anchor="middle" font-size="14" fill="#EDE9FE">Automatic Failover</text>

  <!-- ═══════════════════ ARROWS ═══════════════════ -->
  <!-- Web App → DAB -->
  <line x1="420" y1="450" x2="575" y2="450" stroke="#2563EB" stroke-width="4" marker-end="url(#arrowGreen)"/>
  <rect x="445" y="425" width="90" height="28" rx="6" fill="#EFF6FF" stroke="#BFDBFE" stroke-width="2"/>
  <text x="490" y="445" text-anchor="middle" font-size="14" font-weight="600" fill="#2563EB">REST</text>

  <!-- DAB → Web App (response) -->
  <line x1="575" y1="510" x2="420" y2="510" stroke="#64748B" stroke-width="2" stroke-dasharray="8 4" marker-end="url(#arrowGray)"/>
  <text x="490" y="540" text-anchor="middle" font-size="12" fill="#64748B">JSON Response</text>

  <!-- DAB → Primary (R/W) -->
  <line x1="890" y1="475" x2="1125" y2="420" stroke="#16A34A" stroke-width="4" marker-end="url(#arrowGreen)"/>
  <rect x="945" y="415" width="100" height="36" rx="8" fill="#F0FDF4" stroke="#BBF7D0" stroke-width="2"/>
  <text x="995" y="432" text-anchor="middle" font-size="13" font-weight="700" fill="#16A34A">Read/Write</text>
  <text x="995" y="446" text-anchor="middle" font-size="11" fill="#16A34A">TDS</text>

  <!-- DAB → HA Replica (Read-only) -->
  <line x1="890" y1="565" x2="1000" y2="565" stroke="#F59E0B" stroke-width="4"/>
  <line x1="1000" y1="565" x2="1000" y2="420" stroke="#F59E0B" stroke-width="4"/>
  <line x1="1000" y1="420" x2="1495" y2="420" stroke="#F59E0B" stroke-width="4" marker-end="url(#arrowAmber)"/>
  <rect x="1200" y="390" width="100" height="36" rx="8" fill="#FFFBEB" stroke="#FDE68A" stroke-width="2"/>
  <text x="1250" y="407" text-anchor="middle" font-size="13" font-weight="700" fill="#D97706">Read-Only</text>
  <text x="1250" y="421" text-anchor="middle" font-size="11" fill="#D97706">TDS</text>

  <!-- Footer -->
  <text x="960" y="1030" text-anchor="middle" font-size="14" fill="#94A3B8">ZavaTax Demo · SQL Server 2025 / Azure SQL DB Hyperscale</text>
</svg>'''


def generate_svg_diagram2():
    """Diagram 2: With AI Integration - adds Azure OpenAI with connections from all replicas"""
    return '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1920 1080" font-family="Segoe UI, Arial, sans-serif">
  <defs>
    <filter id="shadow" x="-4%" y="-4%" width="108%" height="112%">
      <feDropShadow dx="3" dy="4" stdDeviation="6" flood-opacity="0.12"/>
    </filter>
    <marker id="arrowGreen" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#16A34A"/>
    </marker>
    <marker id="arrowAmber" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#F59E0B"/>
    </marker>
    <marker id="arrowGray" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#64748B"/>
    </marker>
    <marker id="arrowRed" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#DC2626"/>
    </marker>
  </defs>

  <!-- Background -->
  <rect width="1920" height="1080" fill="#F8FAFC"/>

  <!-- Title -->
  <text x="960" y="55" text-anchor="middle" font-size="36" font-weight="700" fill="#1E293B">ZavaTax Architecture — AI Integration</text>
  <text x="960" y="90" text-anchor="middle" font-size="18" fill="#64748B">Web App + DAB + Hyperscale + Azure OpenAI (Embeddings &amp; Chat via Microsoft Foundry)</text>

  <!-- ═══════════════════ WEB APP ═══════════════════ -->
  <rect x="50" y="200" width="280" height="320" rx="18" fill="#2563EB" filter="url(#shadow)"/>
  <text x="190" y="250" text-anchor="middle" font-size="22" font-weight="700" fill="#FFF">Web Application</text>
  <line x1="80" y1="268" x2="300" y2="268" stroke="#93C5FD" stroke-width="2"/>
  <text x="190" y="300" text-anchor="middle" font-size="15" fill="#DBEAFE">React + TypeScript</text>
  
  <rect x="75" y="330" width="230" height="40" rx="8" fill="rgba(255,255,255,0.15)"/>
  <text x="190" y="356" text-anchor="middle" font-size="14" font-weight="600" fill="#FFF">CQRS Router</text>
  
  <rect x="75" y="385" width="230" height="40" rx="8" fill="rgba(255,255,255,0.15)"/>
  <text x="190" y="411" text-anchor="middle" font-size="14" font-weight="600" fill="#FFF">AI Search &amp; Chat UI</text>
  
  <rect x="75" y="440" width="230" height="55" rx="8" fill="rgba(255,255,255,0.1)"/>
  <text x="190" y="465" text-anchor="middle" font-size="12" fill="#DBEAFE">/ai-assistant · /ai-ask</text>
  <text x="190" y="483" text-anchor="middle" font-size="12" fill="#DBEAFE">/search-cases</text>

  <!-- ═══════════════════ DAB ═══════════════════ -->
  <rect x="420" y="200" width="280" height="320" rx="18" fill="#059669" filter="url(#shadow)"/>
  <text x="560" y="250" text-anchor="middle" font-size="22" font-weight="700" fill="#FFF">Data API Builder</text>
  <line x1="450" y1="268" x2="670" y2="268" stroke="#6EE7B7" stroke-width="2"/>
  <text x="560" y="300" text-anchor="middle" font-size="15" fill="#D1FAE5">REST + GraphQL</text>
  
  <rect x="445" y="325" width="230" height="55" rx="8" fill="rgba(255,255,255,0.2)"/>
  <circle cx="468" cy="352" r="10" fill="#FDE047"/>
  <text x="488" y="345" font-size="13" font-weight="600" fill="#FFF">Primary (R/W)</text>
  <text x="488" y="365" font-size="11" fill="#D1FAE5">CRUD + AI Procs</text>

  <rect x="445" y="395" width="230" height="55" rx="8" fill="rgba(255,255,255,0.2)"/>
  <circle cx="468" cy="422" r="10" fill="#F59E0B"/>
  <text x="488" y="415" font-size="13" font-weight="600" fill="#FFF">Read-Only</text>
  <text x="488" y="435" font-size="11" fill="#D1FAE5">ApplicationIntent=ReadOnly</text>

  <rect x="445" y="465" width="230" height="35" rx="6" fill="rgba(0,0,0,0.15)"/>
  <text x="560" y="488" text-anchor="middle" font-size="12" fill="#FFF">Auto-gen CRUD + Stored Procs</text>

  <!-- ═══════════════════ HYPERSCALE ═══════════════════ -->
  <rect x="780" y="140" width="560" height="420" rx="20" fill="none" stroke="#7C3AED" stroke-width="4" stroke-dasharray="12 6"/>
  <text x="1060" y="175" text-anchor="middle" font-size="20" font-weight="700" fill="#7C3AED">Azure SQL DB Hyperscale</text>

  <!-- Primary Replica -->
  <rect x="810" y="205" width="250" height="160" rx="14" fill="#7C3AED" filter="url(#shadow)"/>
  <text x="935" y="245" text-anchor="middle" font-size="20" font-weight="700" fill="#FFF">Primary Replica</text>
  <text x="935" y="272" text-anchor="middle" font-size="14" fill="#EDE9FE">Read / Write</text>
  <rect x="835" y="290" width="200" height="55" rx="6" fill="rgba(255,255,255,0.15)"/>
  <text x="935" y="312" text-anchor="middle" font-size="11" fill="#EDE9FE">VECTOR · DiskANN</text>
  <text x="935" y="330" text-anchor="middle" font-size="11" fill="#EDE9FE">AI_GENERATE_EMBEDDINGS()</text>

  <!-- HA Replica -->
  <rect x="1080" y="205" width="250" height="160" rx="14" fill="#8B5CF6" filter="url(#shadow)"/>
  <text x="1205" y="245" text-anchor="middle" font-size="20" font-weight="700" fill="#FFF">HA Replica</text>
  <text x="1205" y="272" text-anchor="middle" font-size="14" fill="#EDE9FE">Read-Only</text>
  <rect x="1105" y="290" width="200" height="55" rx="6" fill="rgba(255,255,255,0.15)"/>
  <text x="1205" y="312" text-anchor="middle" font-size="11" fill="#EDE9FE">Zone-Redundant</text>
  <text x="1205" y="330" text-anchor="middle" font-size="11" fill="#EDE9FE">AI_GENERATE_EMBEDDINGS()</text>

  <!-- ═══════════════════ AZURE OPENAI ═══════════════════ -->
  <rect x="780" y="620" width="560" height="280" rx="20" fill="none" stroke="#EA580C" stroke-width="4" stroke-dasharray="12 6"/>
  <text x="1060" y="660" text-anchor="middle" font-size="20" font-weight="700" fill="#EA580C">Azure OpenAI · Microsoft Foundry</text>

  <!-- Embeddings Model -->
  <rect x="810" y="690" width="250" height="130" rx="14" fill="#EA580C" filter="url(#shadow)"/>
  <text x="935" y="730" text-anchor="middle" font-size="18" font-weight="700" fill="#FFF">Embeddings</text>
  <text x="935" y="758" text-anchor="middle" font-size="13" fill="#FFF3E0">text-embedding-3-small</text>
  <text x="935" y="782" text-anchor="middle" font-size="12" fill="#FFF3E0">VECTOR(1536)</text>
  <text x="935" y="805" text-anchor="middle" font-size="11" fill="#FFF3E0">CREATE EXTERNAL MODEL</text>

  <!-- Chat Model -->
  <rect x="1080" y="690" width="250" height="130" rx="14" fill="#EA580C" filter="url(#shadow)"/>
  <text x="1205" y="730" text-anchor="middle" font-size="18" font-weight="700" fill="#FFF">Chat Completions</text>
  <text x="1205" y="758" text-anchor="middle" font-size="13" fill="#FFF3E0">GPT · RAG Responses</text>
  <text x="1205" y="782" text-anchor="middle" font-size="12" fill="#FFF3E0">sp_invoke_external_rest_endpoint</text>
  <text x="1205" y="805" text-anchor="middle" font-size="11" fill="#FFF3E0">Managed Identity Auth</text>

  <!-- ═══════════════════ ARROWS ═══════════════════ -->
  <!-- Web App → DAB -->
  <line x1="330" y1="360" x2="415" y2="360" stroke="#2563EB" stroke-width="4" marker-end="url(#arrowGreen)"/>
  <rect x="340" y="340" width="55" height="22" rx="4" fill="#EFF6FF" stroke="#BFDBFE" stroke-width="1"/>
  <text x="368" y="356" text-anchor="middle" font-size="11" font-weight="600" fill="#2563EB">REST</text>

  <!-- DAB → Primary (R/W) -->
  <line x1="700" y1="352" x2="805" y2="285" stroke="#16A34A" stroke-width="4" marker-end="url(#arrowGreen)"/>
  <rect x="710" y="290" width="70" height="28" rx="5" fill="#F0FDF4" stroke="#BBF7D0" stroke-width="1"/>
  <text x="745" y="308" text-anchor="middle" font-size="11" font-weight="700" fill="#16A34A">R/W</text>

  <!-- DAB → HA Replica (Read-only) -->
  <line x1="700" y1="422" x2="745" y2="422" stroke="#F59E0B" stroke-width="4"/>
  <line x1="745" y1="422" x2="745" y2="285" stroke="#F59E0B" stroke-width="4"/>
  <line x1="745" y1="285" x2="1075" y2="285" stroke="#F59E0B" stroke-width="4" marker-end="url(#arrowAmber)"/>
  <rect x="870" y="260" width="70" height="28" rx="5" fill="#FFFBEB" stroke="#FDE68A" stroke-width="1"/>
  <text x="905" y="278" text-anchor="middle" font-size="11" font-weight="700" fill="#D97706">Read</text>

  <!-- Primary → Embeddings -->
  <line x1="935" y1="365" x2="935" y2="560" stroke="#DC2626" stroke-width="3"/>
  <line x1="935" y1="560" x2="935" y2="685" stroke="#DC2626" stroke-width="3" marker-end="url(#arrowRed)"/>
  
  <!-- Primary → Chat -->
  <line x1="980" y1="365" x2="980" y2="500" stroke="#DC2626" stroke-width="3"/>
  <line x1="980" y1="500" x2="1205" y2="500" stroke="#DC2626" stroke-width="3"/>
  <line x1="1205" y1="500" x2="1205" y2="685" stroke="#DC2626" stroke-width="3" marker-end="url(#arrowRed)"/>

  <!-- HA → Embeddings -->
  <line x1="1130" y1="365" x2="1130" y2="450" stroke="#DC2626" stroke-width="3"/>
  <line x1="1130" y1="450" x2="870" y2="450" stroke="#DC2626" stroke-width="3"/>
  <line x1="870" y1="450" x2="870" y2="580" stroke="#DC2626" stroke-width="3"/>
  <line x1="870" y1="580" x2="870" y2="685" stroke="#DC2626" stroke-width="3" marker-end="url(#arrowRed)"/>

  <!-- HA → Chat -->
  <line x1="1280" y1="365" x2="1280" y2="520" stroke="#DC2626" stroke-width="3"/>
  <line x1="1280" y1="520" x2="1280" y2="685" stroke="#DC2626" stroke-width="3" marker-end="url(#arrowRed)"/>

  <!-- AI Label -->
  <rect x="1350" y="460" width="120" height="60" rx="8" fill="#FEF2F2" stroke="#FECACA" stroke-width="2"/>
  <text x="1410" y="485" text-anchor="middle" font-size="12" font-weight="600" fill="#DC2626">External REST</text>
  <text x="1410" y="505" text-anchor="middle" font-size="11" fill="#DC2626">HTTPS</text>

  <!-- Footer -->
  <text x="960" y="1040" text-anchor="middle" font-size="14" fill="#94A3B8">ZavaTax Demo · SQL Server 2025 / Azure SQL DB Hyperscale · Azure OpenAI</text>
</svg>'''


def generate_svg_diagram3():
    """Diagram 3: Full Architecture - adds Named Replica with analytics routing"""
    return '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1920 1080" font-family="Segoe UI, Arial, sans-serif">
  <defs>
    <filter id="shadow" x="-4%" y="-4%" width="108%" height="112%">
      <feDropShadow dx="3" dy="4" stdDeviation="6" flood-opacity="0.12"/>
    </filter>
    <marker id="arrowGreen" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#16A34A"/>
    </marker>
    <marker id="arrowAmber" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#F59E0B"/>
    </marker>
    <marker id="arrowCyan" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#0891B2"/>
    </marker>
    <marker id="arrowRed" markerWidth="12" markerHeight="8" refX="10" refY="4" orient="auto">
      <polygon points="0 0, 12 4, 0 8" fill="#DC2626"/>
    </marker>
  </defs>

  <!-- Background -->
  <rect width="1920" height="1080" fill="#F8FAFC"/>

  <!-- Title -->
  <text x="960" y="50" text-anchor="middle" font-size="34" font-weight="700" fill="#1E293B">ZavaTax Architecture — Complete</text>
  <text x="960" y="82" text-anchor="middle" font-size="16" fill="#64748B">Web App + DAB (CQRS) + Hyperscale (Primary + HA + Named Replica) + Azure OpenAI</text>

  <!-- ═══════════════════ WEB APP ═══════════════════ -->
  <rect x="40" y="180" width="240" height="360" rx="16" fill="#2563EB" filter="url(#shadow)"/>
  <text x="160" y="225" text-anchor="middle" font-size="20" font-weight="700" fill="#FFF">Web Application</text>
  <line x1="65" y1="242" x2="255" y2="242" stroke="#93C5FD" stroke-width="2"/>
  <text x="160" y="270" text-anchor="middle" font-size="14" fill="#DBEAFE">React + TypeScript</text>
  
  <rect x="60" y="295" width="200" height="36" rx="7" fill="rgba(255,255,255,0.15)"/>
  <text x="160" y="318" text-anchor="middle" font-size="13" font-weight="600" fill="#FFF">CQRS Router</text>
  
  <rect x="60" y="345" width="200" height="80" rx="7" fill="rgba(255,255,255,0.1)"/>
  <text x="160" y="367" text-anchor="middle" font-size="11" fill="#DBEAFE">R/W → Primary</text>
  <text x="160" y="385" text-anchor="middle" font-size="11" fill="#DBEAFE">Read → HA Replica</text>
  <text x="160" y="403" text-anchor="middle" font-size="11" fill="#DBEAFE">Analytics → Named Replica</text>
  <text x="160" y="421" text-anchor="middle" font-size="11" fill="#DBEAFE">AI → Primary → OpenAI</text>

  <rect x="60" y="440" width="200" height="80" rx="7" fill="rgba(255,255,255,0.12)"/>
  <text x="160" y="462" text-anchor="middle" font-size="12" font-weight="600" fill="#FFF">5 Personas</text>
  <text x="160" y="480" text-anchor="middle" font-size="10" fill="#DBEAFE">Tax Filer · Professional</text>
  <text x="160" y="496" text-anchor="middle" font-size="10" fill="#DBEAFE">Manager · Executive · DevOps</text>

  <!-- ═══════════════════ DAB ═══════════════════ -->
  <rect x="350" y="140" width="260" height="430" rx="16" fill="#059669" filter="url(#shadow)"/>
  <text x="480" y="180" text-anchor="middle" font-size="20" font-weight="700" fill="#FFF">Data API Builder</text>
  <line x1="375" y1="195" x2="585" y2="195" stroke="#6EE7B7" stroke-width="2"/>
  <text x="480" y="220" text-anchor="middle" font-size="14" fill="#D1FAE5">REST + GraphQL</text>
  
  <!-- Primary source -->
  <rect x="370" y="245" width="220" height="65" rx="8" fill="rgba(255,255,255,0.2)"/>
  <circle cx="390" cy="277" r="9" fill="#FDE047"/>
  <text x="408" y="268" font-size="12" font-weight="600" fill="#FFF">Primary (R/W)</text>
  <text x="408" y="286" font-size="10" fill="#D1FAE5">CRUD + AI Stored Procs</text>
  <text x="408" y="301" font-size="9" fill="#D1FAE5">/branches · /returns · /ai-assistant</text>

  <!-- Read-only source -->
  <rect x="370" y="325" width="220" height="65" rx="8" fill="rgba(255,255,255,0.2)"/>
  <circle cx="390" cy="357" r="9" fill="#F59E0B"/>
  <text x="408" y="348" font-size="12" font-weight="600" fill="#FFF">Read-Only</text>
  <text x="408" y="366" font-size="10" fill="#D1FAE5">ApplicationIntent=ReadOnly</text>
  <text x="408" y="381" font-size="9" fill="#D1FAE5">/branches-read · /returns-read</text>

  <!-- Analytics source -->
  <rect x="370" y="405" width="220" height="65" rx="8" fill="rgba(255,255,255,0.2)"/>
  <circle cx="390" cy="437" r="9" fill="#0891B2"/>
  <text x="408" y="428" font-size="12" font-weight="600" fill="#FFF">Analytics</text>
  <text x="408" y="446" font-size="10" fill="#D1FAE5">Named Replica FQDN</text>
  <text x="408" y="461" font-size="9" fill="#D1FAE5">/branch-analytics · /executive-kpis</text>

  <rect x="370" y="485" width="220" height="65" rx="8" fill="rgba(0,0,0,0.12)"/>
  <text x="480" y="510" text-anchor="middle" font-size="11" font-weight="600" fill="#FFF">Config Files</text>
  <text x="480" y="526" text-anchor="middle" font-size="9" fill="#D1FAE5">dab-config.json (primary)</text>
  <text x="480" y="540" text-anchor="middle" font-size="9" fill="#D1FAE5">dab-config.replica.json / analytics.json</text>

  <!-- ═══════════════════ HYPERSCALE ═══════════════════ -->
  <rect x="680" y="110" width="600" height="470" rx="18" fill="none" stroke="#7C3AED" stroke-width="4" stroke-dasharray="12 6"/>
  <text x="980" y="145" text-anchor="middle" font-size="18" font-weight="700" fill="#7C3AED">Azure SQL DB Hyperscale</text>

  <!-- Primary Replica -->
  <rect x="705" y="170" width="270" height="140" rx="12" fill="#7C3AED" filter="url(#shadow)"/>
  <text x="840" y="205" text-anchor="middle" font-size="18" font-weight="700" fill="#FFF">Primary Replica</text>
  <text x="840" y="228" text-anchor="middle" font-size="13" fill="#EDE9FE">Read / Write</text>
  <rect x="725" y="242" width="230" height="50" rx="6" fill="rgba(255,255,255,0.15)"/>
  <text x="840" y="262" text-anchor="middle" font-size="10" fill="#EDE9FE">VECTOR · DiskANN · Columnstore · JSON</text>
  <text x="840" y="280" text-anchor="middle" font-size="10" fill="#EDE9FE">AI_GENERATE_EMBEDDINGS()</text>

  <!-- HA Replica -->
  <rect x="995" y="170" width="270" height="140" rx="12" fill="#8B5CF6" filter="url(#shadow)"/>
  <text x="1130" y="205" text-anchor="middle" font-size="18" font-weight="700" fill="#FFF">HA Replica</text>
  <text x="1130" y="228" text-anchor="middle" font-size="13" fill="#EDE9FE">Read-Only · Zone-Redundant</text>
  <rect x="1015" y="242" width="230" height="50" rx="6" fill="rgba(255,255,255,0.15)"/>
  <text x="1130" y="262" text-anchor="middle" font-size="10" fill="#EDE9FE">Automatic Failover</text>
  <text x="1130" y="280" text-anchor="middle" font-size="10" fill="#EDE9FE">AI_GENERATE_EMBEDDINGS()</text>

  <!-- Named Replica -->
  <rect x="850" y="340" width="375" height="160" rx="12" fill="#0891B2" filter="url(#shadow)"/>
  <text x="1037" y="378" text-anchor="middle" font-size="18" font-weight="700" fill="#FFF">Named Replica</text>
  <text x="1037" y="402" text-anchor="middle" font-size="13" fill="#E0F2FE">Read-Only · Dedicated Compute · Separate FQDN</text>
  <rect x="875" y="418" width="325" height="65" rx="6" fill="rgba(255,255,255,0.15)"/>
  <text x="1037" y="438" text-anchor="middle" font-size="10" fill="#E0F2FE">Independent Scaling · Analytics Workloads</text>
  <text x="1037" y="456" text-anchor="middle" font-size="10" fill="#E0F2FE">Columnstore Indexes · Branch Analytics · KPIs</text>
  <text x="1037" y="474" text-anchor="middle" font-size="10" fill="#E0F2FE">AI_GENERATE_EMBEDDINGS()</text>

  <!-- ═══════════════════ AZURE OPENAI ═══════════════════ -->
  <rect x="1320" y="110" width="560" height="470" rx="18" fill="none" stroke="#EA580C" stroke-width="4" stroke-dasharray="12 6"/>
  <text x="1600" y="145" text-anchor="middle" font-size="18" font-weight="700" fill="#EA580C">Azure OpenAI · Microsoft Foundry</text>

  <!-- Embeddings Model -->
  <rect x="1350" y="180" width="230" height="150" rx="12" fill="#EA580C" filter="url(#shadow)"/>
  <text x="1465" y="218" text-anchor="middle" font-size="17" font-weight="700" fill="#FFF">Embeddings</text>
  <text x="1465" y="245" text-anchor="middle" font-size="12" fill="#FFF3E0">text-embedding-3-small</text>
  <text x="1465" y="268" text-anchor="middle" font-size="11" fill="#FFF3E0">VECTOR(1536)</text>
  <rect x="1370" y="282" width="190" height="35" rx="5" fill="rgba(255,255,255,0.15)"/>
  <text x="1465" y="303" text-anchor="middle" font-size="10" fill="#FFF">CREATE EXTERNAL MODEL</text>

  <!-- Chat Model -->
  <rect x="1620" y="180" width="230" height="150" rx="12" fill="#EA580C" filter="url(#shadow)"/>
  <text x="1735" y="218" text-anchor="middle" font-size="17" font-weight="700" fill="#FFF">Chat Completions</text>
  <text x="1735" y="245" text-anchor="middle" font-size="12" fill="#FFF3E0">GPT · RAG Responses</text>
  <text x="1735" y="268" text-anchor="middle" font-size="11" fill="#FFF3E0">Managed Identity</text>
  <rect x="1640" y="282" width="190" height="35" rx="5" fill="rgba(255,255,255,0.15)"/>
  <text x="1735" y="303" text-anchor="middle" font-size="10" fill="#FFF">sp_invoke_external_rest_endpoint</text>

  <!-- Service callout -->
  <rect x="1420" y="420" width="330" height="80" rx="10" fill="#FFF7ED" stroke="#FDBA74" stroke-width="2"/>
  <text x="1585" y="450" text-anchor="middle" font-size="13" font-weight="600" fill="#C2410C">Server-Side AI Calls</text>
  <text x="1585" y="472" text-anchor="middle" font-size="11" fill="#C2410C">All Hyperscale replicas can invoke</text>
  <text x="1585" y="490" text-anchor="middle" font-size="11" fill="#C2410C">external models via HTTPS REST</text>

  <!-- ═══════════════════ ARROWS: Web → DAB ═══════════════════ -->
  <line x1="280" y1="360" x2="345" y2="360" stroke="#2563EB" stroke-width="3" marker-end="url(#arrowGreen)"/>
  <rect x="285" y="343" width="45" height="18" rx="3" fill="#EFF6FF" stroke="#BFDBFE" stroke-width="1"/>
  <text x="308" y="356" text-anchor="middle" font-size="10" font-weight="600" fill="#2563EB">REST</text>

  <!-- ═══════════════════ ARROWS: DAB → Hyperscale ═══════════════════ -->
  <!-- DAB → Primary -->
  <line x1="590" y1="277" x2="700" y2="240" stroke="#16A34A" stroke-width="3" marker-end="url(#arrowGreen)"/>
  <rect x="610" y="232" width="55" height="22" rx="4" fill="#F0FDF4" stroke="#BBF7D0" stroke-width="1"/>
  <text x="638" y="247" text-anchor="middle" font-size="10" font-weight="700" fill="#16A34A">R/W</text>

  <!-- DAB → HA Replica -->
  <line x1="590" y1="357" x2="635" y2="357" stroke="#F59E0B" stroke-width="3"/>
  <line x1="635" y1="357" x2="635" y2="240" stroke="#F59E0B" stroke-width="3"/>
  <line x1="635" y1="240" x2="990" y2="240" stroke="#F59E0B" stroke-width="3" marker-end="url(#arrowAmber)"/>
  <rect x="780" y="218" width="55" height="22" rx="4" fill="#FFFBEB" stroke="#FDE68A" stroke-width="1"/>
  <text x="808" y="233" text-anchor="middle" font-size="10" font-weight="700" fill="#D97706">Read</text>

  <!-- DAB → Named Replica -->
  <line x1="590" y1="437" x2="655" y2="437" stroke="#0891B2" stroke-width="3"/>
  <line x1="655" y1="437" x2="655" y2="420" stroke="#0891B2" stroke-width="3"/>
  <line x1="655" y1="420" x2="845" y2="420" stroke="#0891B2" stroke-width="3" marker-end="url(#arrowCyan)"/>
  <rect x="710" y="398" width="75" height="22" rx="4" fill="#ECFEFF" stroke="#A5F3FC" stroke-width="1"/>
  <text x="748" y="413" text-anchor="middle" font-size="10" font-weight="700" fill="#0891B2">Analytics</text>

  <!-- ═══════════════════ ARROWS: Hyperscale → OpenAI ═══════════════════ -->
  <!-- Primary → Embeddings -->
  <line x1="975" y1="240" x2="1100" y2="240" stroke="#DC2626" stroke-width="2"/>
  <line x1="1100" y1="240" x2="1100" y2="160" stroke="#DC2626" stroke-width="2"/>
  <line x1="1100" y1="160" x2="1465" y2="160" stroke="#DC2626" stroke-width="2"/>
  <line x1="1465" y1="160" x2="1465" y2="175" stroke="#DC2626" stroke-width="2" marker-end="url(#arrowRed)"/>

  <!-- Primary → Chat -->
  <line x1="975" y1="260" x2="1060" y2="260" stroke="#DC2626" stroke-width="2"/>
  <line x1="1060" y1="260" x2="1060" y2="130" stroke="#DC2626" stroke-width="2"/>
  <line x1="1060" y1="130" x2="1735" y2="130" stroke="#DC2626" stroke-width="2"/>
  <line x1="1735" y1="130" x2="1735" y2="175" stroke="#DC2626" stroke-width="2" marker-end="url(#arrowRed)"/>

  <!-- HA → Embeddings -->
  <line x1="1245" y1="310" x2="1245" y2="350" stroke="#DC2626" stroke-width="2"/>
  <line x1="1245" y1="350" x2="1300" y2="350" stroke="#DC2626" stroke-width="2"/>
  <line x1="1300" y1="350" x2="1300" y2="255" stroke="#DC2626" stroke-width="2"/>
  <line x1="1300" y1="255" x2="1345" y2="255" stroke="#DC2626" stroke-width="2" marker-end="url(#arrowRed)"/>

  <!-- HA → Chat -->
  <line x1="1265" y1="240" x2="1285" y2="240" stroke="#DC2626" stroke-width="2"/>
  <line x1="1285" y1="240" x2="1285" y2="370" stroke="#DC2626" stroke-width="2"/>
  <line x1="1285" y1="370" x2="1850" y2="370" stroke="#DC2626" stroke-width="2"/>
  <line x1="1850" y1="370" x2="1850" y2="255" stroke="#DC2626" stroke-width="2"/>
  <line x1="1850" y1="255" x2="1735" y2="255" stroke="#DC2626" stroke-width="2" marker-end="url(#arrowRed)"/>

  <!-- Named → Embeddings -->
  <line x1="1225" y1="420" x2="1310" y2="420" stroke="#DC2626" stroke-width="2"/>
  <line x1="1310" y1="420" x2="1310" y2="330" stroke="#DC2626" stroke-width="2"/>
  <line x1="1310" y1="330" x2="1345" y2="330" stroke="#DC2626" stroke-width="2" marker-end="url(#arrowRed)"/>

  <!-- Named → Chat -->
  <line x1="1225" y1="450" x2="1620" y2="450" stroke="#DC2626" stroke-width="2"/>
  <line x1="1620" y1="450" x2="1620" y2="330" stroke="#DC2626" stroke-width="2" marker-end="url(#arrowRed)"/>

  <!-- Footer -->
  <text x="960" y="1040" text-anchor="middle" font-size="14" fill="#94A3B8">ZavaTax Demo · SQL Server 2025 / Azure SQL DB Hyperscale · Azure OpenAI · Data API Builder</text>
</svg>'''


# =============================================================================
# PPTX Generation
# =============================================================================

def add_rounded_rectangle(slide, left, top, width, height, fill_color, text="", font_size=14, font_color="FFFFFF", bold=False):
    """Add a rounded rectangle shape with text."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb(fill_color)
    shape.line.fill.background()  # No border
    
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.color.rgb = hex_to_rgb(font_color)
        run.font.bold = bold
        tf.paragraphs[0].space_before = Pt(0)
        tf.paragraphs[0].space_after = Pt(0)
    
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=12, font_color="1E293B", bold=False, align=PP_ALIGN.CENTER):
    """Add a text box."""
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = hex_to_rgb(font_color)
    run.font.bold = bold
    return shape


def add_dashed_rectangle(slide, left, top, width, height, border_color):
    """Add a dashed rectangle (no fill)."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.background()  # No fill
    shape.line.color.rgb = hex_to_rgb(border_color)
    shape.line.width = Pt(3)
    shape.line.dash_style = 4  # Dash style
    return shape


def add_arrow(slide, start_x, start_y, end_x, end_y, color):
    """Add a connector arrow."""
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(start_x), Inches(start_y),
        Inches(end_x), Inches(end_y)
    )
    connector.line.color.rgb = hex_to_rgb(color)
    connector.line.width = Pt(3)
    return connector


def add_circle(slide, left, top, size, fill_color):
    """Add a circle."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(left), Inches(top), Inches(size), Inches(size)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb(fill_color)
    shape.line.fill.background()
    return shape


def create_slide1(prs):
    """Slide 1: Base Architecture"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Title
    add_text_box(slide, 0, 0.2, 13.33, 0.5, "ZavaTax Architecture — Base", 28, COLORS['text_dark'], True)
    add_text_box(slide, 0, 0.65, 13.33, 0.35, "Web Application + Data API Builder + Azure SQL DB Hyperscale", 14, '64748B')

    # Web Application
    add_rounded_rectangle(slide, 0.5, 1.5, 2.5, 3.2, COLORS['web_app'], "Web Application", 18, COLORS['text_light'], True)
    add_text_box(slide, 0.55, 2.1, 2.4, 0.3, "React + TypeScript", 12, COLORS['text_light'])
    add_rounded_rectangle(slide, 0.7, 2.6, 2.1, 0.4, '1D4ED8', "CQRS Router", 11, COLORS['text_light'], True)
    add_text_box(slide, 0.55, 3.2, 2.4, 0.8, "R/W → Primary\nRead → HA Replica", 10, 'DBEAFE')

    # DAB
    add_rounded_rectangle(slide, 3.8, 1.5, 2.5, 3.2, COLORS['dab'], "Data API Builder", 18, COLORS['text_light'], True)
    add_text_box(slide, 3.85, 2.1, 2.4, 0.3, "REST + GraphQL", 12, COLORS['text_light'])
    
    # DAB data sources
    add_rounded_rectangle(slide, 4.0, 2.5, 2.1, 0.5, '047857', "", 10)
    add_circle(slide, 4.1, 2.62, 0.25, 'FDE047')
    add_text_box(slide, 4.4, 2.55, 1.6, 0.4, "Primary (R/W)", 10, COLORS['text_light'], True)
    
    add_rounded_rectangle(slide, 4.0, 3.15, 2.1, 0.5, '047857', "", 10)
    add_circle(slide, 4.1, 3.27, 0.25, 'F59E0B')
    add_text_box(slide, 4.4, 3.2, 1.6, 0.4, "Read-Only", 10, COLORS['text_light'], True)

    # Hyperscale boundary
    add_dashed_rectangle(slide, 7.0, 1.2, 5.8, 3.8, COLORS['hyperscale_border'])
    add_text_box(slide, 7.2, 1.35, 5.4, 0.35, "Azure SQL DB Hyperscale", 16, COLORS['hyperscale_border'], True)

    # Primary Replica
    add_rounded_rectangle(slide, 7.3, 1.9, 2.5, 1.4, COLORS['primary'], "Primary Replica", 16, COLORS['text_light'], True)
    add_text_box(slide, 7.35, 2.4, 2.4, 0.3, "Read / Write", 11, 'EDE9FE')
    add_rounded_rectangle(slide, 7.5, 2.75, 2.1, 0.4, '6D28D9', "VECTOR · DiskANN", 9, 'EDE9FE')

    # HA Replica
    add_rounded_rectangle(slide, 10.1, 1.9, 2.5, 1.4, COLORS['ha_replica'], "HA Replica", 16, COLORS['text_light'], True)
    add_text_box(slide, 10.15, 2.4, 2.4, 0.3, "Read-Only", 11, 'EDE9FE')
    add_rounded_rectangle(slide, 10.3, 2.75, 2.1, 0.4, '7C3AED', "Zone-Redundant", 9, 'EDE9FE')

    # Arrows (simplified - using shapes as connectors can be tricky)
    # Web → DAB
    add_rounded_rectangle(slide, 3.1, 2.8, 0.6, 0.25, '2563EB', "REST", 9, COLORS['text_light'], True)
    
    # DAB → Primary label
    add_rounded_rectangle(slide, 6.35, 2.2, 0.55, 0.35, '16A34A', "R/W", 9, COLORS['text_light'], True)
    
    # DAB → HA label
    add_rounded_rectangle(slide, 6.35, 3.0, 0.55, 0.35, 'F59E0B', "Read", 9, COLORS['text_light'], True)

    return slide


def create_slide2(prs):
    """Slide 2: With AI Integration"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Title
    add_text_box(slide, 0, 0.15, 13.33, 0.45, "ZavaTax Architecture — AI Integration", 26, COLORS['text_dark'], True)
    add_text_box(slide, 0, 0.55, 13.33, 0.3, "Web App + DAB + Hyperscale + Azure OpenAI (Embeddings & Chat)", 13, '64748B')

    # Web Application (smaller)
    add_rounded_rectangle(slide, 0.3, 1.2, 2.0, 2.4, COLORS['web_app'], "Web App", 15, COLORS['text_light'], True)
    add_text_box(slide, 0.35, 1.65, 1.9, 0.25, "React + TypeScript", 10, COLORS['text_light'])
    add_rounded_rectangle(slide, 0.45, 2.0, 1.7, 0.35, '1D4ED8', "CQRS Router", 9, COLORS['text_light'], True)
    add_rounded_rectangle(slide, 0.45, 2.5, 1.7, 0.35, '1D4ED8', "AI Search UI", 9, COLORS['text_light'], True)

    # DAB (smaller)
    add_rounded_rectangle(slide, 2.8, 1.2, 2.0, 2.4, COLORS['dab'], "Data API Builder", 14, COLORS['text_light'], True)
    add_text_box(slide, 2.85, 1.65, 1.9, 0.25, "REST + GraphQL", 10, COLORS['text_light'])
    
    add_rounded_rectangle(slide, 2.95, 2.0, 1.7, 0.35, '047857', "", 9)
    add_circle(slide, 3.0, 2.08, 0.18, 'FDE047')
    add_text_box(slide, 3.2, 2.05, 1.4, 0.25, "Primary", 9, COLORS['text_light'], True)
    
    add_rounded_rectangle(slide, 2.95, 2.5, 1.7, 0.35, '047857', "", 9)
    add_circle(slide, 3.0, 2.58, 0.18, 'F59E0B')
    add_text_box(slide, 3.2, 2.55, 1.4, 0.25, "Read-Only", 9, COLORS['text_light'], True)

    # Hyperscale boundary
    add_dashed_rectangle(slide, 5.3, 1.0, 3.8, 2.8, COLORS['hyperscale_border'])
    add_text_box(slide, 5.4, 1.1, 3.6, 0.3, "Azure SQL DB Hyperscale", 12, COLORS['hyperscale_border'], True)

    # Primary Replica
    add_rounded_rectangle(slide, 5.5, 1.5, 1.7, 1.0, COLORS['primary'], "Primary", 13, COLORS['text_light'], True)
    add_text_box(slide, 5.55, 1.9, 1.6, 0.2, "R/W", 9, 'EDE9FE')
    add_text_box(slide, 5.55, 2.15, 1.6, 0.25, "AI_GENERATE_EMBEDDINGS()", 7, 'EDE9FE')

    # HA Replica
    add_rounded_rectangle(slide, 7.4, 1.5, 1.6, 1.0, COLORS['ha_replica'], "HA Replica", 12, COLORS['text_light'], True)
    add_text_box(slide, 7.45, 1.9, 1.5, 0.2, "Read-Only", 9, 'EDE9FE')
    add_text_box(slide, 7.45, 2.15, 1.5, 0.25, "AI_GENERATE_EMBEDDINGS()", 7, 'EDE9FE')

    # Azure OpenAI boundary
    add_dashed_rectangle(slide, 5.3, 4.1, 3.8, 2.4, COLORS['ai_border'])
    add_text_box(slide, 5.4, 4.2, 3.6, 0.3, "Azure OpenAI · Microsoft Foundry", 12, COLORS['ai_border'], True)

    # Embeddings
    add_rounded_rectangle(slide, 5.5, 4.6, 1.7, 1.0, COLORS['azure_ai'], "Embeddings", 12, COLORS['text_light'], True)
    add_text_box(slide, 5.55, 5.0, 1.6, 0.2, "text-embedding-3-small", 8, 'FFF3E0')
    add_text_box(slide, 5.55, 5.25, 1.6, 0.2, "VECTOR(1536)", 8, 'FFF3E0')

    # Chat
    add_rounded_rectangle(slide, 7.4, 4.6, 1.6, 1.0, COLORS['azure_ai'], "Chat", 12, COLORS['text_light'], True)
    add_text_box(slide, 7.45, 5.0, 1.5, 0.2, "GPT · RAG", 8, 'FFF3E0')
    add_text_box(slide, 7.45, 5.25, 1.5, 0.2, "Managed Identity", 8, 'FFF3E0')

    # Arrow labels
    add_rounded_rectangle(slide, 2.35, 2.0, 0.4, 0.2, '2563EB', "REST", 8, COLORS['text_light'])
    add_rounded_rectangle(slide, 4.85, 1.7, 0.4, 0.2, '16A34A', "R/W", 8, COLORS['text_light'])
    add_rounded_rectangle(slide, 4.85, 2.3, 0.4, 0.2, 'F59E0B', "Read", 8, COLORS['text_light'])
    
    # AI call labels
    add_rounded_rectangle(slide, 6.1, 3.6, 1.4, 0.35, 'FEF2F2', "External REST (HTTPS)", 8, 'DC2626')

    return slide


def create_slide3(prs):
    """Slide 3: Full Architecture with Named Replica"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Title
    add_text_box(slide, 0, 0.1, 13.33, 0.4, "ZavaTax Architecture — Complete", 24, COLORS['text_dark'], True)
    add_text_box(slide, 0, 0.45, 13.33, 0.25, "Web App + DAB (CQRS) + Hyperscale (Primary + HA + Named Replica) + Azure OpenAI", 11, '64748B')

    # Web Application
    add_rounded_rectangle(slide, 0.2, 0.9, 1.7, 2.6, COLORS['web_app'], "Web App", 13, COLORS['text_light'], True)
    add_text_box(slide, 0.25, 1.3, 1.6, 0.2, "React + TS", 9, COLORS['text_light'])
    add_rounded_rectangle(slide, 0.3, 1.55, 1.5, 0.28, '1D4ED8', "CQRS Router", 8, COLORS['text_light'], True)
    add_text_box(slide, 0.25, 1.95, 1.6, 0.7, "R/W → Primary\nRead → HA\nAnalytics → Named\n5 Personas", 7, 'DBEAFE')

    # DAB
    add_rounded_rectangle(slide, 2.3, 0.9, 1.8, 2.6, COLORS['dab'], "Data API Builder", 12, COLORS['text_light'], True)
    add_text_box(slide, 2.35, 1.3, 1.7, 0.2, "REST + GraphQL", 9, COLORS['text_light'])
    
    add_rounded_rectangle(slide, 2.4, 1.55, 1.6, 0.35, '047857')
    add_circle(slide, 2.45, 1.63, 0.15, 'FDE047')
    add_text_box(slide, 2.62, 1.58, 1.3, 0.25, "Primary (R/W)", 8, COLORS['text_light'])
    
    add_rounded_rectangle(slide, 2.4, 2.0, 1.6, 0.35, '047857')
    add_circle(slide, 2.45, 2.08, 0.15, 'F59E0B')
    add_text_box(slide, 2.62, 2.03, 1.3, 0.25, "Read-Only", 8, COLORS['text_light'])
    
    add_rounded_rectangle(slide, 2.4, 2.45, 1.6, 0.35, '047857')
    add_circle(slide, 2.45, 2.53, 0.15, '0891B2')
    add_text_box(slide, 2.62, 2.48, 1.3, 0.25, "Analytics", 8, COLORS['text_light'])

    # Hyperscale boundary
    add_dashed_rectangle(slide, 4.5, 0.75, 4.3, 2.9, COLORS['hyperscale_border'])
    add_text_box(slide, 4.6, 0.85, 4.1, 0.25, "Azure SQL DB Hyperscale", 11, COLORS['hyperscale_border'], True)

    # Primary
    add_rounded_rectangle(slide, 4.65, 1.2, 1.95, 0.9, COLORS['primary'], "Primary", 11, COLORS['text_light'], True)
    add_text_box(slide, 4.7, 1.55, 1.85, 0.2, "R/W · AI_GENERATE_EMBEDDINGS()", 7, 'EDE9FE')

    # HA
    add_rounded_rectangle(slide, 6.75, 1.2, 1.95, 0.9, COLORS['ha_replica'], "HA Replica", 11, COLORS['text_light'], True)
    add_text_box(slide, 6.8, 1.55, 1.85, 0.2, "Read · Zone-Redundant", 7, 'EDE9FE')
    add_text_box(slide, 6.8, 1.75, 1.85, 0.2, "AI_GENERATE_EMBEDDINGS()", 7, 'EDE9FE')

    # Named Replica
    add_rounded_rectangle(slide, 4.65, 2.3, 4.05, 1.15, COLORS['named_replica'], "Named Replica", 12, COLORS['text_light'], True)
    add_text_box(slide, 4.7, 2.7, 3.95, 0.2, "Read-Only · Dedicated Compute · Separate FQDN", 8, 'E0F2FE')
    add_text_box(slide, 4.7, 2.95, 3.95, 0.2, "Analytics · Columnstore · AI_GENERATE_EMBEDDINGS()", 8, 'E0F2FE')

    # Azure OpenAI boundary
    add_dashed_rectangle(slide, 9.2, 0.75, 3.9, 2.9, COLORS['ai_border'])
    add_text_box(slide, 9.3, 0.85, 3.7, 0.25, "Azure OpenAI · Microsoft Foundry", 11, COLORS['ai_border'], True)

    # Embeddings
    add_rounded_rectangle(slide, 9.4, 1.2, 1.7, 0.9, COLORS['azure_ai'], "Embeddings", 11, COLORS['text_light'], True)
    add_text_box(slide, 9.45, 1.55, 1.6, 0.18, "text-embedding-3-small", 7, 'FFF3E0')
    add_text_box(slide, 9.45, 1.78, 1.6, 0.18, "VECTOR(1536)", 7, 'FFF3E0')

    # Chat
    add_rounded_rectangle(slide, 11.3, 1.2, 1.6, 0.9, COLORS['azure_ai'], "Chat", 11, COLORS['text_light'], True)
    add_text_box(slide, 11.35, 1.55, 1.5, 0.18, "GPT · RAG", 8, 'FFF3E0')
    add_text_box(slide, 11.35, 1.78, 1.5, 0.18, "Managed Identity", 7, 'FFF3E0')

    # Callout
    add_rounded_rectangle(slide, 9.5, 2.5, 3.4, 0.8, 'FFF7ED', "", 9)
    add_text_box(slide, 9.55, 2.6, 3.3, 0.6, "All Hyperscale replicas can invoke\nexternal models via HTTPS REST", 9, 'C2410C')

    # Arrow labels
    add_rounded_rectangle(slide, 1.95, 1.9, 0.3, 0.18, '2563EB', "REST", 7, COLORS['text_light'])
    add_rounded_rectangle(slide, 4.15, 1.35, 0.3, 0.18, '16A34A', "R/W", 7, COLORS['text_light'])
    add_rounded_rectangle(slide, 4.15, 1.6, 0.3, 0.18, 'F59E0B', "Read", 7, COLORS['text_light'])
    add_rounded_rectangle(slide, 4.15, 2.6, 0.35, 0.18, '0891B2', "Anlytcs", 6, COLORS['text_light'])

    return slide


def generate_pptx():
    """Generate the complete PPTX with all 3 slides."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 aspect ratio
    prs.slide_height = Inches(7.5)
    
    create_slide1(prs)
    create_slide2(prs)
    create_slide3(prs)
    
    return prs


# =============================================================================
# Main
# =============================================================================

def main():
    output_dir = os.path.dirname(os.path.abspath(__file__))
    
    # Generate SVGs
    print("Generating SVG diagrams...")
    
    svg1_path = os.path.join(output_dir, "01_base_architecture.svg")
    with open(svg1_path, "w", encoding="utf-8") as f:
        f.write(generate_svg_diagram1())
    print(f"  ✓ {svg1_path}")
    
    svg2_path = os.path.join(output_dir, "02_ai_integration.svg")
    with open(svg2_path, "w", encoding="utf-8") as f:
        f.write(generate_svg_diagram2())
    print(f"  ✓ {svg2_path}")
    
    svg3_path = os.path.join(output_dir, "03_full_architecture.svg")
    with open(svg3_path, "w", encoding="utf-8") as f:
        f.write(generate_svg_diagram3())
    print(f"  ✓ {svg3_path}")
    
    # Generate PPTX
    print("\nGenerating PPTX...")
    pptx_path = os.path.join(output_dir, "ZavaTax_Architecture.pptx")
    prs = generate_pptx()
    prs.save(pptx_path)
    print(f"  ✓ {pptx_path}")
    
    print("\n✅ All diagrams generated successfully!")
    print(f"   Output directory: {output_dir}")


if __name__ == "__main__":
    main()
